# SPDX-License-Identifier: GPL-3.0-or-later
# Copyright (C) 2026 Marco Sumari
# This file is part of IngePresupuestos — https://ingepresupuestos.com
# Software libre bajo la GNU GPL v3 o posterior. Ver el archivo LICENSE.
"""Issue #14 — franjas de imagen de la empresa en cabecera y pie.

Una empresa que evaluaba el programa necesitaba su membrete y su tira de
logos en los presupuestos. Decisión de Marco (21-10-2026): cada franja
REEMPLAZA a su banda de texto; en el pie, el número de página sigue, encima
de la franja. Sale en el PDF (y con él el Reporte completo y Control de
Obra), el Gantt, la Curva S y Word; en Excel solo la de cabecera.

Corre con:  QT_QPA_PLATFORM=offscreen venv/bin/python3 tests/test_franjas_imagen.py

Usa una copia temporal del seed, nunca la BD activa.
"""
import io
import os
import shutil
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))
os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')

from PySide6.QtCore import QBuffer, QByteArray, QIODevice
from PySide6.QtGui import QColor, QImage
from PySide6.QtWidgets import QApplication

_app = QApplication.instance() or QApplication([])

import core.config as cfg
import core.database as d

SEED = os.path.join(os.path.dirname(__file__), '..', 'presupuestos_seed.db')
_fd, _tmpdb = tempfile.mkstemp(suffix='_franjas_test.db')
os.close(_fd)
shutil.copy(SEED, _tmpdb)
d.DB_PATH = _tmpdb
cfg.DB_PATH = _tmpdb
d.init_db()

import core.pdf_reports as pr   # noqa: E402  (después de fijar la BD)

PID = 397
EMPRESA = 'CONSTRUCTORA DE PRUEBA SAC'
_conn = d.get_db()
_conn.execute("UPDATE proyectos SET cliente='MUNICIPALIDAD DE PRUEBA' WHERE id=?", (PID,))
_conn.commit()
_conn.close()


def _png_b64(w: int, h: int, color: str) -> str:
    img = QImage(w, h, QImage.Format_RGB32)
    img.fill(QColor(color))
    ba = QByteArray()
    buf = QBuffer(ba)
    buf.open(QIODevice.WriteOnly)
    img.save(buf, 'PNG')
    return bytes(ba.toBase64()).decode('ascii')


CABECERA = _png_b64(1600, 200, '#0B3D91')     # 8:1, entra entera a lo ancho
PIE = _png_b64(1600, 100, '#00695C')


def _formato(cabecera='', pie='', **banderas):
    d.set_config('rep_empresa_nombre', EMPRESA)
    d.set_config('rep_cabecera_img_b64', cabecera)
    d.set_config('rep_pie_img_b64', pie)
    for k in ('rep_encabezado_oculto', 'rep_pie_oculto', 'rep_pie_der_oculto'):
        d.set_config(k, '1' if banderas.get(k) else '0')


def _tmp(sufijo):
    fd, path = tempfile.mkstemp(suffix=sufijo)
    os.close(fd)
    return path


def _pagina_pdf(**kw):
    """Palabras de la primera página del Presupuesto sin portada."""
    import pdfplumber
    _formato(**kw)
    path = _tmp('_franjas.pdf')
    try:
        pr.generar_pdf_archivo('presupuesto', PID, path, with_cover=False)
        with pdfplumber.open(path) as pdf:
            pg = pdf.pages[0]
            palabras = pg.extract_words()
            return ' '.join(w['text'] for w in palabras), palabras, pg.height
    finally:
        os.unlink(path)


# ── Las reglas comunes ───────────────────────────────────────────────────────

def test_sin_franjas_la_geometria_es_la_de_siempre():
    _formato()
    r = pr._PdfRenderer({'nombre': 'x'}, 'x')
    assert r.franja_cab is None and r.franja_pie is None
    assert r.margin_top_body == int(1.05 * r.dpi)
    assert r.margin_bot_body == int(round(18 / 25.4 * r.dpi))


def test_apagar_el_encabezado_o_el_pie_manda_sobre_la_franja():
    _formato(CABECERA, PIE, rep_encabezado_oculto=True, rep_pie_oculto=True)
    assert pr.franja_cabecera() is None and pr.franja_pie() is None
    _formato(CABECERA, PIE)
    assert pr.franja_cabecera() is not None and pr.franja_pie() is not None


def test_la_franja_llena_el_ancho_sin_deformarse_y_respeta_el_tope():
    img = QImage(1600, 200, QImage.Format_RGB32)          # 8:1
    x, w, h = pr.rect_franja(img, 50, 800, 300)
    assert (x, w, h) == (50, 800, 100)
    img = QImage(400, 400, QImage.Format_RGB32)           # cuadrada: la frena el tope
    x, w, h = pr.rect_franja(img, 50, 800, 100)
    assert (w, h) == (100, 100) and x == 50 + 350         # centrada


# ── PDF ──────────────────────────────────────────────────────────────────────

def test_el_pdf_cambia_el_encabezado_de_texto_por_la_franja():
    texto, palabras, _alto = _pagina_pdf()
    assert 'CONSTRUCTORA' in texto and 'Costo' in texto
    texto, palabras, _alto = _pagina_pdf(cabecera=CABECERA)
    assert 'CONSTRUCTORA' not in texto and 'Costo' not in texto
    # El cuerpo arranca debajo de la franja.
    r = pr._PdfRenderer({'nombre': 'x'}, 'x')
    pt = 72.0 / r.dpi
    tops = [w['top'] for w in palabras if w['text'] == 'Descripción']
    assert tops and min(tops) > r.franja_cab_rect.bottom() * pt


def test_el_pie_con_franja_deja_solo_el_numero_de_pagina():
    texto, _p, _a = _pagina_pdf()
    assert 'Cliente:' in texto
    texto, palabras, alto = _pagina_pdf(pie=PIE)
    assert 'Cliente:' not in texto
    r = pr._PdfRenderer({'nombre': 'x'}, 'x')
    pt = 72.0 / r.dpi
    pagina = [w for w in palabras if w['text'] == 'Página']
    assert pagina, 'falta el número de página'
    # encima de la franja, no sobre ella
    assert pagina[0]['bottom'] <= r.franja_pie_rect.top() * pt + 0.5
    # y si el usuario lo apagó, tampoco sale
    texto, _p, _a = _pagina_pdf(pie=PIE, rep_pie_der_oculto=True)
    assert 'Página' not in texto


# ── Gantt ────────────────────────────────────────────────────────────────────

def test_el_gantt_cambia_el_encabezado_por_la_franja():
    import pypdf
    from views.cronograma_view import CronogramaView
    conn = d.get_db()
    proy = dict(conn.execute("SELECT * FROM proyectos WHERE id=?", (PID,)).fetchone())
    conn.close()
    cv = CronogramaView(PID, proy, lambda: None)
    cv.show()
    try:
        cv.cargar()
    except Exception:
        pass

    def _texto():
        path = _tmp('_gantt.pdf')
        try:
            cv._gantt_w._render_pdf_completo(path, 'multi', 'landscape', True, papel='A4')
            return ' '.join(pypdf.PdfReader(path).pages[0].extract_text().split())
        finally:
            os.unlink(path)

    _formato()
    sin = _texto()
    assert 'CONSTRUCTORA' in sin and 'MUNICIPALIDAD' in sin
    _formato(CABECERA, PIE)
    con = _texto()
    assert 'CONSTRUCTORA' not in con and 'MUNICIPALIDAD' not in con
    assert 'Página 1 de' in con
    cv.close()


# ── Word y Excel ─────────────────────────────────────────────────────────────

def test_word_lleva_las_franjas_en_lugar_de_las_tablas():
    import docx
    import core.word_reports as WR
    path = _tmp('_franjas.docx')
    try:
        _formato(CABECERA, PIE)
        WR.generar_word('resumen', PID, path)
        sec = docx.Document(path).sections[0]
        assert not sec.header.tables and not sec.footer.tables
        assert sec.header._element.xpath('.//pic:pic')
        assert sec.footer._element.xpath('.//pic:pic')
        assert ' PAGE ' in sec.footer._element.xml     # el número de página sigue
        _formato()
        WR.generar_word('resumen', PID, path)
        sec = docx.Document(path).sections[0]
        assert sec.header.tables and sec.footer.tables
    finally:
        os.unlink(path)


def test_excel_lleva_la_franja_de_cabecera_en_a1():
    import openpyxl
    import core.exporter as EX
    _formato(CABECERA)
    ws = openpyxl.load_workbook(io.BytesIO(EX.exportar_presupuesto(PID).getvalue())).worksheets[0]
    assert len(ws._images) == 1
    fila1 = [str(c.value or '') for c in ws[1]]
    assert not any('Costo al' in v for v in fila1)
    _formato()
    ws = openpyxl.load_workbook(io.BytesIO(EX.exportar_presupuesto(PID).getvalue())).worksheets[0]
    assert any('Costo al' in str(c.value or '') for c in ws[1])


# ── «Editar formato» ─────────────────────────────────────────────────────────

def test_editar_formato_guarda_y_quita_las_franjas():
    from views.formato_reporte_dialog import FormatoReporteDialog
    _formato()
    dlg = FormatoReporteDialog()
    dlg._poner_franja('rep_cabecera_img_b64', CABECERA)
    dlg._poner_franja('rep_pie_img_b64', PIE)
    lbl, btn_quitar = dlg._franja_widgets['rep_pie_img_b64']
    assert btn_quitar.isEnabled() and not lbl.pixmap().isNull()
    dlg._save_and_accept()
    f = pr.get_formato()
    assert f['rep_cabecera_img_b64'] == CABECERA and f['rep_pie_img_b64'] == PIE
    dlg = FormatoReporteDialog()
    dlg._franja_widgets['rep_cabecera_img_b64'][1].click()     # «Quitar»
    dlg._save_and_accept()
    assert pr.get_formato()['rep_cabecera_img_b64'] == ''
    _formato()


if __name__ == '__main__':
    import pytest
    sys.exit(pytest.main([__file__, '-q']))
