# SPDX-License-Identifier: GPL-3.0-or-later
# Copyright (C) 2026 Marco Sumari
# This file is part of IngePresupuestos — https://ingepresupuestos.com
# Software libre bajo la GNU GPL v3 o posterior. Ver el archivo LICENSE.
"""Issue #31 (parte de la 3.0.18): el ACU impreso con aire.

Corre con:  QT_QPA_PLATFORM=offscreen venv/bin/python3 -m pytest tests/test_acu_aire.py

Cada tipo de recurso va en su bloque con un rótulo (MANO DE OBRA,
MATERIALES…) y un espacio entre bloques, en el PDF y en el Excel; los
subtotales siguen juntos al final. Usa una COPIA temporal del seed.
"""
import os
import shutil
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))
os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')

from PySide6.QtWidgets import QApplication                          # noqa: E402

import core.config as cfg                                           # noqa: E402
import core.database as d                                           # noqa: E402

_app = QApplication.instance() or QApplication([])

SEED = os.path.join(os.path.dirname(__file__), '..', 'presupuestos_seed.db')
_fd, _tmpdb = tempfile.mkstemp(suffix='_acu_aire_test.db')
os.close(_fd)
shutil.copy(SEED, _tmpdb)
d.DB_PATH = _tmpdb
cfg.DB_PATH = _tmpdb
d.init_db()

PID = 397  # Yanque; su 01.01 (cartel) tiene 2 MO + 5 MAT + 1 EQ


def _acu_0101(html: str) -> str:
    ini = html.index('01.01')
    return html[ini:html.index('01.02', ini)]


def test_pdf_un_rotulo_por_tipo_en_orden_y_subtotales_al_final():
    from core import pdf_reports
    proy = dict(d.get_db().execute("SELECT * FROM proyectos WHERE id=?",
                                   (PID,)).fetchone())
    items, _tot = d.calcular_totales(PID)
    html = _acu_0101(pdf_reports._html_acus(PID, proy, items))
    pos = [html.index(f'>{r}</td>') for r in ('MANO DE OBRA', 'MATERIALES', 'EQUIPOS')]
    assert pos == sorted(pos)
    assert 'SUB-CONTRATOS' not in html          # sin recursos SC, sin rótulo
    assert html.count('class="acu-grupo"') == 3
    # Los recursos de cada tipo quedan bajo su rótulo
    assert pos[0] < html.index('OPERARIO') < pos[1] < html.index('CEMENTO') < pos[2]
    # Subtotales juntos al final, después del último recurso
    assert html.index('HERRAMIENTAS') < html.index('>Mano de obra</td>') \
        < html.index('COSTO UNITARIO TOTAL')


def test_excel_un_rotulo_por_tipo_antes_de_sus_filas():
    from core.exporter import exportar_acus
    import openpyxl
    wb = openpyxl.load_workbook(exportar_acus(PID))
    ws = wb.active
    col_a = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
    ini = next(i for i, v in enumerate(col_a) if v and str(v).startswith('01.01'))
    fin = next(i for i, v in enumerate(col_a) if v and str(v).startswith('01.02'))
    bloque = col_a[ini:fin]
    i_mo, i_mat, i_eq = (bloque.index(r) for r in ('MANO DE OBRA', 'MATERIALES', 'EQUIPOS'))
    assert bloque[i_mo + 1:i_mo + 3] == ['MO', 'MO']
    assert bloque[i_mat + 1:i_mat + 6] == ['MAT'] * 5
    assert bloque[i_eq + 1] == 'EQ'
    # Una fila vacía de aire antes de cada rótulo que no es el primero
    assert bloque[i_mat - 1] is None and bloque[i_eq - 1] is None
    # Subtotales juntos al final
    assert bloque.index('Mano de obra') > i_eq + 1
    assert bloque.index('COSTO UNITARIO TOTAL') > bloque.index('Equipos')


if __name__ == '__main__':
    import pytest
    sys.exit(pytest.main([__file__, '-q']))
