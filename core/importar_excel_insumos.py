# SPDX-License-Identifier: GPL-3.0-or-later
# Copyright (C) 2026 Marco Sumari
# This file is part of IngePresupuestos — https://ingepresupuestos.com
# Software libre bajo la GNU GPL v3 o posterior. Ver el archivo LICENSE.
"""Importar insumos desde un Excel cualquiera, asignando columnas.

Hasta octubre de 2026 «Importar ▸ Desde Excel» del Catálogo de Insumos solo
aceptaba el formato que la propia app exporta: la PRIMERA fila con Código ·
Tipo · Descripción · Unidad · Precio. Una lista oficial con su propio formato
no entraba. Lo pidieron a la vez un usuario de Uruguay (la Lista Oficial de
Precios del MTOP, que sale cada mes en Excel), uno de España (su base de
precios exportada desde Menfis) y los de Brasil (SINAPI).

Este módulo es la lógica pura —sin Qt— que usa
`views/importar_excel_dialog.py`:

* `leer_libro` — todas las hojas a memoria (.xlsx con openpyxl; .xls con xlrd
  si está instalado).
* `detectar_cabecera` / `auto_asignar` — la fila de títulos y qué columna es
  qué, por nombre, tolerando tildes, mayúsculas y sinónimos en español,
  portugués e inglés.
* `detectar_cantidad_en_unidad` / `detectar_impuesto` — las dos rarezas de
  las listas oficiales: «100 Kg» en la columna de unidad (el precio es de las
  100) y el precio CON IVA.
* `es_lista_mtop` / `configuracion_inicial` — la Lista del MTOP queda
  configurada sola.
* `preparar_insumos` → `importar_insumos` — filas → insumos → BD (actualiza
  por código, si no crea; un 0 nunca pisa un precio existente).
"""
from __future__ import annotations

import hashlib
import json
import re
import unicodedata
from dataclasses import asdict, dataclass, field

from core.config import INEI_DEFAULT, TIPOS_RECURSO
from utils.formatting import parse_num, parse_num_opt

# Campos de destino, en el orden en que se muestran.
CAMPOS = ('codigo', 'descripcion', 'unidad', 'precio', 'tipo', 'inei')
OBLIGATORIOS = ('descripcion', 'precio')

# Filas que se miran para encontrar la cabecera. Las listas oficiales llevan
# dos o tres filas de título encima; nadie pone treinta.
FILAS_BUSQUEDA_CABECERA = 30

# Motivos de omisión (clave → texto del resumen).
MOTIVOS = {
    'sin_precio':      "sin precio (0 o vacío)",
    'sin_descripcion': "sin descripción",
    'tipo_invalido':   "tipo no reconocido (MO/MAT/EQ/SC)",
    'texto':           "notas o subtítulos (sin código ni precio)",
}


# ── Normalización de textos de cabecera ─────────────────────────────────────
def normalizar(s) -> str:
    """«PRECIOS MAS IVA» → «precios mas iva»; «Descrição» → «descricao»;
    «N°» → «n»; «c/IVA» → «c iva». Sin tildes, sin puntuación, en minúsculas
    y con un solo espacio entre palabras."""
    if s is None:
        return ''
    t = unicodedata.normalize('NFKD', str(s).lower())
    t = ''.join(c for c in t if not unicodedata.combining(c))
    t = re.sub(r'[^a-z0-9]+', ' ', t)
    return t.strip()


# Sinónimos por campo, ya normalizados. Un sinónimo de varias palabras se
# busca como frase dentro de la cabecera; uno de una palabra, como palabra
# entera («und» no debe encajar en «fundación»).
SINONIMOS: dict[str, tuple[str, ...]] = {
    'descripcion': ('descripcion', 'denominacion', 'descricao', 'description',
                    'designacion', 'nombre', 'concepto', 'detalle', 'resumen',
                    'insumo', 'recurso', 'material', 'materiales', 'texto'),
    'precio':      ('precio', 'precios', 'preco', 'precos', 'valor', 'importe',
                    'costo', 'coste', 'custo', 'price', 'pu',
                    'precio unitario', 'costo unitario'),
    'unidad':      ('unidad', 'unidades', 'und', 'unid', 'ud', 'u', 'um',
                    'unidade', 'unit', 'medida', 'cantidad'),
    'codigo':      ('codigo', 'cod', 'code', 'nro', 'n', 'no', 'num', 'numero',
                    'item', 'clave', 'ref', 'referencia', 'id'),
    'tipo':        ('tipo', 'clase', 'type', 'naturaleza'),
    'inei':        ('indice inei', 'inei', 'indice unificado', 'iu'),
}

# Orden de asignación: primero lo obligatorio, y «tipo» antes que «inei»
# para que «Índice» suelto no se lleve otra cosa.
_ORDEN_ASIGNACION = ('descripcion', 'precio', 'unidad', 'codigo', 'tipo', 'inei')

_TIPOS_SINONIMOS = {
    'MO':  ('mo', 'mano de obra', 'mano obra', 'mao de obra', 'labor', 'personal'),
    'MAT': ('mat', 'mt', 'material', 'materiales', 'materia'),
    'EQ':  ('eq', 'equipo', 'equipos', 'herramienta', 'herramientas',
            'maquinaria', 'equipamento'),
    'SC':  ('sc', 'subcontrato', 'subcontratos', 'sub contrato', 'servicio',
            'servicios', 'servico'),
}


def _puntaje(campo: str, cabecera_norm: str) -> int:
    """2 = la cabecera ES un sinónimo; 1 = lo contiene; 0 = nada."""
    if not cabecera_norm:
        return 0
    sin = SINONIMOS[campo]
    if cabecera_norm in sin:
        return 2
    palabras = cabecera_norm.split()
    for s in sin:
        if ' ' in s:
            if f' {s} ' in f' {cabecera_norm} ':
                return 1
        elif s in palabras:
            return 1
    return 0


def campo_tipo(valor) -> str | None:
    """Texto de una celda de «Tipo» → MO/MAT/EQ/SC, o None si no se entiende."""
    n = normalizar(valor)
    if not n:
        return None
    if n.upper() in TIPOS_RECURSO:
        return n.upper()
    for tipo, sins in _TIPOS_SINONIMOS.items():
        if n in sins:
            return tipo
    return None


# ── Lectura del libro ───────────────────────────────────────────────────────
@dataclass
class Libro:
    ruta: str
    hojas: list[str]
    filas: dict[str, list[list]]     # hoja → filas (listas de valores)
    activa: str = ''


def _recortar(filas: list[list]) -> list[list]:
    """Quita columnas vacías a la derecha y filas vacías al final.

    La hoja «NUEVOS MATERIALES» del MTOP declara 1 944 filas y usa 30: lo de
    abajo es formato sin datos."""
    out = []
    for f in filas:
        f = list(f)
        while f and (f[-1] is None or str(f[-1]).strip() == ''):
            f.pop()
        out.append(f)
    while out and not out[-1]:
        out.pop()
    return out


def leer_libro(ruta: str, max_filas: int = 50000) -> Libro:
    """Lee todas las hojas. Lanza la excepción del lector si el archivo no se
    puede abrir (quien llama la muestra)."""
    if ruta.lower().endswith('.xls'):
        import xlrd            # opcional: solo para el .xls viejo
        wb = xlrd.open_workbook(ruta)
        hojas, filas = [], {}
        for sh in wb.sheets():
            hojas.append(sh.name)
            filas[sh.name] = _recortar(
                [sh.row_values(i) for i in range(min(sh.nrows, max_filas))])
        return Libro(ruta, hojas, filas, hojas[0] if hojas else '')

    from openpyxl import load_workbook
    wb = load_workbook(ruta, read_only=True, data_only=True)
    try:
        hojas, filas = [], {}
        for ws in wb.worksheets:
            hojas.append(ws.title)
            rows = []
            for i, r in enumerate(ws.iter_rows(values_only=True)):
                if i >= max_filas:
                    break
                rows.append(list(r))
            filas[ws.title] = _recortar(rows)
        activa = wb.active.title if wb.active is not None else (hojas[0] if hojas else '')
    finally:
        wb.close()
    return Libro(ruta, hojas, filas, activa)


# ── Cabecera y asignación ───────────────────────────────────────────────────
def _es_texto_titulo(v) -> bool:
    """Celda que parece un título de columna: texto corto con alguna letra
    («+», «2026» o «1.234,56» no lo son; «Precio junio 2026» sí)."""
    if v is None or isinstance(v, (int, float)):
        return False
    s = str(v).strip()
    return bool(s) and len(s) <= 40 and re.search(r'[^\W\d_]', s) is not None


def detectar_cabecera(filas: list[list]) -> int:
    """Índice (base 0) de la fila de títulos.

    La primera fila donde DOS celdas distintas suenan a un campo conocido
    («DESCRIPCIÓN», «CANTIDAD», «PRECIOS MAS IVA»). Un título suelto de una
    sola celda —«LISTA OFICIAL DE PRECIOS» trae «precios»— no basta. Si
    ninguna fila tiene vocabulario conocido, la primera con tres o más textos
    cortos; si tampoco, la primera."""
    tope = min(len(filas), FILAS_BUSQUEDA_CABECERA)
    for i in range(tope):
        n = sum(1 for v in filas[i]
                if _es_texto_titulo(v)
                and any(_puntaje(c, normalizar(v)) for c in CAMPOS))
        if n >= 2:
            return i
    for i in range(tope):
        if sum(1 for v in filas[i] if _es_texto_titulo(v)) >= 3:
            return i
    return 0


def datos_de(filas: list[list], fila_cabecera: int) -> list[list]:
    return filas[fila_cabecera + 1:]


def _celda(fila: list, col) -> object:
    if col is None or col < 0 or col >= len(fila):
        return None
    return fila[col]


def _texto(v) -> str:
    """Valor de celda → texto. 4001.0 (Excel guarda enteros como float) → «4001»."""
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def auto_asignar(cabecera: list, datos: list[list]) -> dict[str, int | None]:
    """{campo: índice de columna o None}, por el nombre de cada columna.

    Cada columna se usa una sola vez. Si «código» no aparece por nombre, se
    prueba con la columna SIN título textual («+», «#», vacía) más cercana a
    la izquierda de la descripción cuyos valores están casi todos llenos y no
    se repiten — es la de la Lista del MTOP, titulada «+»."""
    normas = [normalizar(h) if isinstance(h, str) else '' for h in cabecera]
    asign: dict[str, int | None] = {c: None for c in CAMPOS}
    usadas: set[int] = set()
    for campo in _ORDEN_ASIGNACION:
        mejor, mejor_p = None, 0
        for i, h in enumerate(normas):
            if i in usadas:
                continue
            p = _puntaje(campo, h)
            if p > mejor_p:
                mejor, mejor_p = i, p
        if mejor is not None:
            asign[campo] = mejor
            usadas.add(mejor)

    d = asign['descripcion']
    if asign['codigo'] is None and d is not None:
        muestra = [f for f in datos[:80] if any(_texto(v) for v in f)]
        for col in range(d - 1, -1, -1):
            if col in usadas:
                continue
            h = cabecera[col] if col < len(cabecera) else None
            if re.search(r'[A-Za-z]', str(h or '')):
                continue                # tiene título y no es de código
            vals = [_texto(_celda(f, col)) for f in muestra]
            llenos = [v for v in vals if v]
            if (muestra and len(llenos) >= 0.8 * len(muestra)
                    and len(set(llenos)) >= 0.9 * len(llenos)
                    and all(len(v) <= 20 for v in llenos)):
                asign['codigo'] = col
                break
    return asign


_RE_CANT_UNIDAD = re.compile(r'^\s*(\d[\d.,]*)\s*([^\d\s].*?)\s*$')


def partir_cantidad_unidad(texto) -> tuple[float, str]:
    """«100 Kg» → (100, «Kg»); «1 m².» → (1, «m²»); «Kg» → (1, «Kg»).

    El punto final de la unidad se quita («m².», «hh.»): es la misma unidad.
    Una cantidad 0 o ilegible cuenta como 1 (no se divide por cero)."""
    s = _texto(texto)
    m = _RE_CANT_UNIDAD.match(s)
    if not m:
        return 1.0, s.rstrip('.').strip()
    cant = parse_num(m.group(1)) or 1.0
    return cant, m.group(2).rstrip('.').strip()


def detectar_cantidad_en_unidad(datos: list[list], col: int | None) -> bool:
    """True si la MAYORÍA de las celdas llenas de la columna Unidad traen
    número + unidad («100 Kg», «5  Kg», «1 l»). Una unidad normal («m2»,
    «kg», «hh») empieza con letra y no cuenta."""
    if col is None:
        return False
    vals = [_texto(_celda(f, col)) for f in datos[:300]]
    vals = [v for v in vals if v]
    if not vals:
        return False
    con = sum(1 for v in vals if _RE_CANT_UNIDAD.match(v))
    return con > len(vals) / 2


_IMPUESTOS = ('iva', 'igv', 'impuesto', 'impuestos', 'itbms', 'itbis', 'isv',
              'icms', 'vat', 'tax')


def detectar_impuesto(cabecera_precio) -> bool:
    """True si el título de la columna de precio dice que lo incluye:
    «PRECIOS MAS IVA», «Precio c/IVA», «Precio con IGV». «sin IVA» y «s/IVA»
    dicen lo contrario."""
    palabras = normalizar(cabecera_precio).split()
    for i, p in enumerate(palabras):
        if p in _IMPUESTOS:
            antes = palabras[i - 1] if i else ''
            return antes not in ('sin', 's', 'excl', 'excluido', 'menos')
    return False


def es_lista_mtop(filas: list[list]) -> bool:
    """La Lista Oficial de Precios del MTOP (Uruguay): lo dice en el título."""
    for f in filas[:6]:
        for v in f:
            if 'lista oficial de precios' in normalizar(v):
                return True
    return False


def firma(cabecera: list) -> str:
    """Huella de un formato: los títulos normalizados, en orden. Sirve para
    recordar la asignación — la Lista del MTOP de julio tiene la misma firma
    que la de junio aunque la hoja se llame distinto."""
    return '|'.join(normalizar(h) for h in cabecera)


def clave_firma(cabecera: list) -> str:
    return hashlib.sha1(firma(cabecera).encode('utf-8')).hexdigest()[:16]


def formato_propio(cabecera: list) -> bool:
    """La cabecera que exporta la propia app (Código·Tipo·Descripción·Unidad·
    Precio[·Índice INEI]). Ahí un precio 0 es un dato, no «sin precio este
    mes» — los recursos en % MO valen 0 —, así que no se omite por defecto."""
    n = [normalizar(h) for h in cabecera if normalizar(h)]
    return n[:5] == ['codigo', 'tipo', 'descripcion', 'unidad', 'precio']


# ── Opciones ────────────────────────────────────────────────────────────────
@dataclass
class Opciones:
    hoja: str = ''
    fila_cabecera: int = 0                 # base 0
    columnas: dict = field(default_factory=lambda: {c: None for c in CAMPOS})
    tipo_fijo: str | None = None           # 'MAT' para toda una lista de materiales
    cantidad_en_unidad: bool = False
    quitar_impuesto: bool = False
    impuesto_pct: float = 18.0
    omitir_sin_precio: bool = True
    prefijo: str = ''
    # Por qué se propuso cada cosa (para explicárselo al usuario).
    motivos: dict = field(default_factory=dict)

    def a_json(self) -> str:
        d = asdict(self)
        d.pop('motivos', None)
        d.pop('hoja', None)
        return json.dumps(d)

    def aplicar_json(self, txt: str) -> bool:
        try:
            d = json.loads(txt)
        except (TypeError, ValueError):
            return False
        cols = d.get('columnas') or {}
        self.columnas = {c: (cols.get(c) if isinstance(cols.get(c), int) else None)
                         for c in CAMPOS}
        self.tipo_fijo = d.get('tipo_fijo') if d.get('tipo_fijo') in TIPOS_RECURSO else None
        for k in ('cantidad_en_unidad', 'quitar_impuesto', 'omitir_sin_precio'):
            if k in d:
                setattr(self, k, bool(d[k]))
        if 'impuesto_pct' in d:
            self.impuesto_pct = parse_num(d['impuesto_pct'])
        self.prefijo = str(d.get('prefijo') or '')
        return True


def configuracion_inicial(filas: list[list], hoja: str, *,
                          impuesto_pct: float,
                          fila_cabecera: int | None = None) -> Opciones:
    """Lo que se propone al abrir una hoja, sin memoria previa."""
    fc = detectar_cabecera(filas) if fila_cabecera is None else fila_cabecera
    cab = filas[fc] if fc < len(filas) else []
    datos = datos_de(filas, fc)
    op = Opciones(hoja=hoja, fila_cabecera=fc, impuesto_pct=impuesto_pct)
    op.columnas = auto_asignar(cab, datos)
    if op.columnas['tipo'] is None:
        op.tipo_fijo = 'MAT'           # sin columna de tipo, lo normal es una lista de materiales
    op.cantidad_en_unidad = detectar_cantidad_en_unidad(datos, op.columnas['unidad'])
    if op.cantidad_en_unidad:
        op.motivos['cantidad'] = ("La columna de unidad trae la cantidad "
                                  "(«100 Kg»): el precio se divide por ella.")
    cp = op.columnas['precio']
    if cp is not None and detectar_impuesto(_celda(cab, cp)):
        op.quitar_impuesto = True
        op.motivos['impuesto'] = (f"El título de la columna de precio "
                                  f"(«{_texto(_celda(cab, cp))}») dice que incluye el impuesto.")
    if formato_propio(cab):
        op.omitir_sin_precio = False
    if es_lista_mtop(filas):
        from core.paises import perfil
        op.tipo_fijo = 'MAT'
        op.columnas['tipo'] = None
        op.columnas['inei'] = None
        op.quitar_impuesto = True
        op.impuesto_pct = float(perfil('UY')['imp'][1])
        op.omitir_sin_precio = True
        op.motivos['mtop'] = True
        op.motivos['impuesto'] = (f"La Lista Oficial del MTOP publica precios "
                                  f"con IVA ({op.impuesto_pct:g} % en Uruguay).")
    return op


def hoja_inicial(libro: Libro) -> str:
    """La hoja que se abre primero: la del MTOP con materiales del mes, o la
    activa si tiene descripción y precio, o la primera que los tenga."""
    def sirve(h):
        f = libro.filas.get(h) or []
        a = auto_asignar(f[detectar_cabecera(f)] if f else [], [])
        return a['descripcion'] is not None and a['precio'] is not None

    mtop = [h for h in libro.hojas if es_lista_mtop(libro.filas.get(h) or [])
            and sirve(h)]
    if mtop:
        normales = [h for h in mtop if 'nuevo' not in normalizar(h)]
        return (normales or mtop)[0]
    if libro.activa and sirve(libro.activa):
        return libro.activa
    for h in libro.hojas:
        if sirve(h):
            return h
    return libro.activa or (libro.hojas[0] if libro.hojas else '')


# ── Filas → insumos ─────────────────────────────────────────────────────────
@dataclass
class Insumo:
    fila: int                  # fila del Excel (base 1), para el resumen
    codigo: str
    tipo: str
    descripcion: str
    unidad: str
    precio: float
    inei: str | None           # None = no vino del archivo (se deduce al guardar)
    original: str = ''         # «100 Kg a 26,222.59» cuando hubo conversión


def faltantes(op: Opciones) -> list[str]:
    """Campos obligatorios sin asignar."""
    return [c for c in OBLIGATORIOS if op.columnas.get(c) is None]


def preparar_insumos(filas: list[list], op: Opciones):
    """→ (insumos, omitidos). `omitidos` = lista de (fila, motivo)."""
    cols = op.columnas
    insumos: list[Insumo] = []
    omitidos: list[tuple[int, str]] = []
    divisor_imp = 1 + (op.impuesto_pct or 0) / 100 if op.quitar_impuesto else 1.0
    for i, f in enumerate(datos_de(filas, op.fila_cabecera)):
        nfila = op.fila_cabecera + 2 + i
        if not any(_texto(v) for v in f):
            continue                    # fila vacía: ni se cuenta
        desc = _texto(_celda(f, cols.get('descripcion')))
        if not desc:
            omitidos.append((nfila, 'sin_descripcion'))
            continue
        pre_raw = _celda(f, cols.get('precio'))
        cod = _texto(_celda(f, cols.get('codigo')))
        # Sin código y con la celda de precio VACÍA (no en 0) es una nota al
        # pie o un subtítulo: «Valores válidos desde el 1/06/2026», «N/C — NO
        # SE COTIZA». Contarla como «sin precio» inflaba ese número.
        if not cod and _texto(pre_raw) == '':
            omitidos.append((nfila, 'texto'))
            continue
        if op.tipo_fijo:
            tipo = op.tipo_fijo
        else:
            tipo = campo_tipo(_celda(f, cols.get('tipo')))
            if tipo is None:
                omitidos.append((nfila, 'tipo_invalido'))
                continue
        und_raw = _celda(f, cols.get('unidad'))
        if op.cantidad_en_unidad:
            cant, unidad = partir_cantidad_unidad(und_raw)
        else:
            cant, unidad = 1.0, _texto(und_raw)
        pre = parse_num_opt(pre_raw) if not isinstance(pre_raw, (int, float)) else float(pre_raw)
        pre = pre or 0.0
        # Un recurso en % (overhead de MO/EQ) vale 0 por definición: no es
        # «sin precio».
        if pre <= 0 and op.omitir_sin_precio and not unidad.startswith('%'):
            omitidos.append((nfila, 'sin_precio'))
            continue
        precio = pre / cant / divisor_imp
        original = ''
        if cant != 1 or divisor_imp != 1:
            original = f"{_texto(und_raw)} a {pre:,.2f}"
            if divisor_imp != 1:
                original += f" (con {op.impuesto_pct:g} %)"
        if cod and op.prefijo:
            cod = op.prefijo + cod
        inei_raw = _texto(_celda(f, cols.get('inei')))
        inei = (inei_raw.zfill(2)[:2] if inei_raw.isdigit() else None)
        insumos.append(Insumo(nfila, cod, tipo, desc, unidad,
                              round(precio, 4), inei, original))
    return insumos, omitidos


def _inei_para(ins: Insumo) -> str:
    """El índice INEI (Perú) cuando el archivo no lo trae: el prefijo de un
    código S10 de 7 dígitos (IU(2) + correlativo(5)), como hacía el
    importador viejo; si el código es de otro sistema («4348», «MTOP-4348»,
    «SA001»), el índice por defecto del tipo. Fuera del Perú no se usa, pero
    la columna es obligatoria en la app."""
    if ins.inei:
        return ins.inei
    if len(ins.codigo) == 7 and ins.codigo.isdigit():
        return ins.codigo[:2]
    return INEI_DEFAULT.get(ins.tipo, '39')


def importar_insumos(insumos: list[Insumo], conn=None) -> dict:
    """Escribe en `recursos`: actualiza por código exacto si existe, si no
    crea. Un precio 0 del archivo NUNCA pisa un precio existente (una lista
    oficial pone 0 cuando ese mes no relevó el insumo). Un insumo existente
    conserva su índice INEI si el archivo no trae columna de índice."""
    from core.database import get_db, _siguiente_codigo_inei
    propia = conn is None
    if propia:
        conn = get_db()
    res = {'nuevos': 0, 'actualizados': 0, 'precio_conservado': 0}
    try:
        for ins in insumos:
            ex = None
            if ins.codigo:
                ex = conn.execute(
                    "SELECT id, precio, indice_inei FROM recursos WHERE codigo=?",
                    (ins.codigo,)).fetchone()
            if ex:
                precio = ins.precio
                if precio <= 0 and (ex['precio'] or 0) > 0:
                    precio = ex['precio']
                    res['precio_conservado'] += 1
                inei = ins.inei or ex['indice_inei'] or _inei_para(ins)
                conn.execute(
                    "UPDATE recursos SET descripcion=?, tipo=?, unidad=?, precio=?, "
                    "indice_inei=? WHERE id=?",
                    (ins.descripcion, ins.tipo, ins.unidad, precio, inei, ex['id']))
                res['actualizados'] += 1
            else:
                inei = _inei_para(ins)
                cod = ins.codigo or _siguiente_codigo_inei(conn, inei)
                conn.execute(
                    "INSERT INTO recursos (codigo, descripcion, tipo, unidad, precio, "
                    "indice_inei) VALUES (?,?,?,?,?,?)",
                    (cod, ins.descripcion, ins.tipo, ins.unidad, ins.precio, inei))
                res['nuevos'] += 1
        conn.commit()
    finally:
        if propia:
            conn.close()
    return res


def resumen(res: dict, omitidos: list[tuple[int, str]]) -> str:
    """Texto final para el usuario: nuevos / actualizados / omitidos y por qué."""
    from utils.i18n import tr
    lineas = [tr("{n} nuevos").format(n=res.get('nuevos', 0)),
              tr("{n} actualizados").format(n=res.get('actualizados', 0))]
    if res.get('precio_conservado'):
        lineas.append(tr("{n} conservaron su precio (el archivo traía 0)")
                      .format(n=res['precio_conservado']))
    if omitidos:
        cuenta: dict[str, list[int]] = {}
        for fila, motivo in omitidos:
            cuenta.setdefault(motivo, []).append(fila)
        lineas.append(tr("{n} filas omitidas:").format(n=len(omitidos)))
        for motivo, fs in cuenta.items():
            ej = ", ".join(str(x) for x in fs[:5]) + ("…" if len(fs) > 5 else "")
            lineas.append(f"   • {len(fs)} {tr(MOTIVOS.get(motivo, motivo))} "
                          f"({tr('filas')} {ej})")
    return "\n".join(lineas)
