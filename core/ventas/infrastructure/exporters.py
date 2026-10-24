# SPDX-License-Identifier: GPL-3.0-or-later
"""Render only supplied document snapshots, never query catalog or construction."""
from contextlib import contextmanager
from html import escape
import os
from pathlib import Path
from tempfile import NamedTemporaryFile

from ..domain.errors import SalesError, ValidationError
from .repositories import fingerprint


def check_document(document):
    if 'totals' not in document or len(document['lines']) != len(document['totals']['lines']):
        raise ValidationError('Calcule/guarde la revisión antes de visualizarla.')
    if document.get('revision_state') == 'EMITIDA':
        if document.get('hash') != fingerprint(document['header'], document['lines'], document['totals']):
            raise ValidationError('Sello de revisión inválido; exportación bloqueada.')


def rows(document):
    check_document(document)
    return [[str(index + 1), line['description'], line['quantity'], line['unit'], line['price'],
             result['line_discount'], result['global_discount'], result['net'],
             result['tax'] if result['tax_known'] else 'SIN DEFINIR', result['total']]
            for index, (line, result) in enumerate(zip(document['lines'], document['totals']['lines']))]


COLUMNS = ['N.º', 'Descripción', 'Cantidad', 'Unidad', 'P. unitario', 'Dto. línea', 'Dto. global', 'Neto', 'Impuesto', 'Total']


def html_preview(document):
    h, t = document['header'], document['totals']
    esc = lambda value: escape(str(value)).replace('\n', '<br/>')
    table_rows = ''.join('<tr>' + ''.join(f'<td>{esc(v)}</td>' for v in row) + '</tr>' for row in rows(document))
    taxes = ''.join(f"<p>{esc(g['name'])} {esc(g['rate'])}% ({esc(g['mode'])}): base {esc(g['base'])} · impuesto {esc(g['tax'])}</p>" for g in t['tax_groups'])
    customer = h.get('customer', {})
    customer_details = ' · '.join(str(customer.get(key, '')) for key in ('document', 'contact', 'email', 'phone') if customer.get(key))
    return f"""<html><body><h2>{esc(h.get('issuer', '') or 'EMISOR SIN CONFIGURAR')}</h2>
    <h3>Cotización {esc(document['number'])} · V{document['version']} · {esc(document.get('revision_state', 'BORRADOR'))}</h3>
    <p>{esc(h['title'])}<br/>Fecha: {esc(h['date'])} · Vigencia: {esc(h['valid_until'])} · Moneda: {esc(h['currency'])}</p>
    <p>Cliente: {esc(customer.get('name', ''))}<br/>{esc(customer_details)}<br/>{esc(customer.get('address', ''))}</p>
    <table border="1" cellpadding="5" cellspacing="0"><tr>{''.join(f'<th>{c}</th>' for c in COLUMNS)}</tr>{table_rows}</table>
    <p>Bruto: {esc(t['gross'])} · Descuentos línea: {esc(t['line_discount'])} · Descuento global: {esc(t['global_discount'])}</p>
    <p>Neto: {esc(t['net'])} · Impuestos: {esc(t['tax']) if t['tax_configured'] else 'SIN CONFIGURAR — TOTAL PROVISIONAL'}</p>
    {taxes}<h3>Total {esc(h['currency'])}: {esc(t['total'])}</h3>
    <p>Condiciones: {esc(h.get('conditions', ''))}</p><p>Observaciones: {esc(h.get('notes', ''))}</p>
    <p>Documento comercial informativo. No es factura ni orden. No reserva ni descuenta inventario.</p>
    <p>Política: {esc(t['policy_version'])} · Sello: {esc(document.get('hash') or 'Borrador sin sello')}</p></body></html>"""


@contextmanager
def atomic_destination(path):
    target = Path(path)
    temporary = None
    try:
        with NamedTemporaryFile(prefix='.ventas-', suffix=target.suffix, dir=target.parent, delete=False) as stream:
            temporary = Path(stream.name)
        yield temporary
        os.replace(temporary, target)
    except OSError as exc:
        raise SalesError('No se pudo escribir el documento. Revise permisos/espacio y cierre el archivo si está abierto.') from exc
    finally:
        if temporary is not None:
            try:
                temporary.unlink(missing_ok=True)
            except OSError:
                pass


def export_pdf(document, path, font_dir):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    check_document(document)
    font_path = Path(font_dir) / 'Inter.ttf'
    if not font_path.is_file():
        raise SalesError('Falta la fuente Inter.ttf para exportación internacional.')
    if 'VentasInter' not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont('VentasInter', str(font_path)))
    style = ParagraphStyle('Ventas', fontName='VentasInter', fontSize=8, leading=11)
    heading = ParagraphStyle('VentasHeading', parent=style, fontSize=14, leading=18, spaceAfter=8)
    p = lambda value: Paragraph(escape(str(value)).replace('\n', '<br/>'), style)
    h, totals = document['header'], document['totals']
    content = [Paragraph(escape(h.get('issuer') or 'EMISOR SIN CONFIGURAR'), heading),
               p(f"Cotización {document['number']} · V{document['version']} · {document['revision_state']}"),
               p(h['title']), p(f"Fecha: {h['date']} · Vigencia: {h['valid_until']} · {h['currency']}"),
               p('Cliente: ' + h['customer'].get('name', '')),
               p(' · '.join(str(h['customer'].get(k, '')) for k in ('document', 'contact', 'phone', 'email') if h['customer'].get(k))),
               p(h['customer'].get('address', '')), Spacer(1, 5 * mm)]
    data = [[p(v) for v in COLUMNS]] + [[p(v) for v in row] for row in rows(document)]
    table = Table(data, colWidths=[9*mm, 68*mm, 23*mm, 16*mm, 26*mm, 24*mm, 24*mm, 26*mm, 26*mm, 28*mm], repeatRows=1)
    table.setStyle(TableStyle([('BACKGROUND', (0,0), (-1,0), colors.HexColor('#F0F1F2')),
                               ('GRID', (0,0), (-1,-1), .25, colors.HexColor('#ABACAE')),
                               ('VALIGN', (0,0), (-1,-1), 'TOP'), ('LEFTPADDING', (0,0), (-1,-1), 4), ('RIGHTPADDING', (0,0), (-1,-1), 4)]))
    content.extend([table, Spacer(1, 4 * mm), p(f"Bruto {totals['gross']} · Dto. línea {totals['line_discount']} · Dto. global {totals['global_discount']}"),
                    p(f"Neto {totals['net']} · Impuestos {totals['tax'] if totals['tax_configured'] else 'SIN CONFIGURAR — PROVISIONAL'}")])
    for group in totals['tax_groups']:
        content.append(p(f"{group['name']} {group['rate']}% {group['mode']} · Base {group['base']} · Impuesto {group['tax']}"))
    content.extend([Paragraph(escape(f"TOTAL {h['currency']}: {totals['total']}"), heading),
                    p('Condiciones: ' + h.get('conditions', '')), p('Observaciones: ' + h.get('notes', '')),
                    p('Documento comercial informativo; no es factura ni orden. No reserva ni descuenta stock.'),
                    p('Política: ' + totals['policy_version']), p('Sello: ' + (document.get('hash') or 'Borrador sin sello'))])
    def footer(canvas, doc):
        canvas.setFont('VentasInter', 8)
        canvas.drawString(10 * mm, 6 * mm, f"{document['number']} · V{document['version']} · Página {doc.page}")
    try:
        with atomic_destination(path) as temp:
            SimpleDocTemplate(str(temp), pagesize=landscape(A4), leftMargin=10*mm, rightMargin=10*mm, topMargin=10*mm, bottomMargin=14*mm).build(content, onFirstPage=footer, onLaterPages=footer)
    except SalesError:
        raise
    except Exception as exc:
        raise SalesError('No se pudo renderizar el PDF de esta revisión. No se reemplazó el archivo destino.') from exc


def export_xlsx(document, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    check_document(document)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Cotización'
    h, totals = document['header'], document['totals']
    data = [[h.get('issuer', '')], ['Cotización', document['number'], 'Versión', str(document['version']), document['revision_state']],
            ['Título', h['title']], ['Cliente', h['customer'].get('name', '')],
            ['Documento', h['customer'].get('document', ''), 'Contacto', h['customer'].get('contact', '')],
            ['Email', h['customer'].get('email', ''), 'Teléfono', h['customer'].get('phone', '')],
            ['Dirección', h['customer'].get('address', '')], ['Fecha', h['date'], 'Vigencia', h['valid_until'], 'Moneda', h['currency']], [], COLUMNS]
    data.extend(rows(document))
    data.extend([[], ['Bruto', totals['gross']], ['Descuento línea', totals['line_discount']], ['Descuento global', totals['global_discount']],
                 ['Neto', totals['net']], ['Impuestos', totals['tax'] if totals['tax_configured'] else 'SIN CONFIGURAR — PROVISIONAL'], ['TOTAL', totals['total']]])
    for group in totals['tax_groups']:
        data.append([group['name'], group['rate'], group['mode'], 'Base', group['base'], 'Impuesto', group['tax']])
    data.extend([['Condiciones', h.get('conditions', '')], ['Observaciones', h.get('notes', '')], ['Sello', document.get('hash') or 'Borrador sin sello'],
                 ['Política', totals['policy_version']], ['Decimales exactos como texto; sin conversión a float ni fórmulas Excel.'],
                 ['Documento comercial informativo; no es factura, orden ni movimiento de stock.']])
    for row_index, values in enumerate(data, 1):
        for column_index, value in enumerate(values, 1):
            cell = sheet.cell(row_index, column_index, str(value))
            # User data starting '=' never becomes an executable spreadsheet formula.
            cell.data_type = 's'
    for cell in sheet[10]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='F0F1F2')
    for index in range(1, len(COLUMNS) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 52 if index == 2 else 21
    sheet.freeze_panes = 'C11'
    try:
        with atomic_destination(path) as temp:
            workbook.save(temp)
    except SalesError:
        raise
    except Exception as exc:
        raise SalesError('No se pudo generar Excel; revise caracteres y longitud de los textos. El archivo destino se conserva.') from exc
    finally:
        workbook.close()
